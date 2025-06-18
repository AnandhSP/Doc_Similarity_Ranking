import os
from docx import Document
from transformers.pipelines import pipeline

class CommunicationAgent:
    def __init__(self, model_name='gpt2'):
        """Initialize the communication agent with a text generation model"""
        self.text_generator = pipeline('text-generation', model=model_name, max_length=200)
    
    def extract_candidate_name(self, profile_path):
        """Extract candidate name from their profile document"""
        try:
            doc = Document(profile_path)
            # Assume name is in the first few paragraphs
            for para in doc.paragraphs[:5]:
                text = para.text.strip()
                if text and len(text.split()) <= 4:  # Most names are 1-4 words
                    return text
            # Fallback to filename
            return os.path.basename(profile_path).replace('.docx', '').replace('_', ' ').title()
        except:
            return os.path.basename(profile_path).replace('.docx', '').replace('_', ' ').title()

    def compose_shortlist_email(self, candidate_info, position_title="the position"):
        """
        Compose a personalized email for a shortlisted candidate
        
        Args:
            candidate_info (dict): Dictionary containing candidate information
                - name: Candidate's name
                - profile_path: Path to candidate's profile
                - similarity_score: Matching score from comparison agent
            position_title (str): Title of the position
        """
        name = candidate_info.get('name', self.extract_candidate_name(candidate_info['profile_path']))
        
        # Email template with placeholders
        email_template = f"""Dear {name},

I hope this email finds you well. I am writing to inform you that after careful review of your profile, we are pleased to announce that you have been shortlisted for {position_title}.

Your application demonstrated strong potential alignment with our requirements, and we would like to proceed with the next steps of our selection process.

You will receive a detailed email shortly with:
- Next round details
- Schedule options
- Required preparations
- Additional documentation (if any)

Please ensure to monitor your email inbox, including the spam folder, for our follow-up communication.

Best regards,
Hexawar Technologies Hiring Team"""

        return email_template

    def process_shortlisted_candidates(self, ranked_results, profile_paths, position_title="the position"):
        """
        Process shortlisted candidates and generate their notification emails
        
        Args:
            ranked_results (list): List of dictionaries with candidate rankings
            profile_paths (list): List of paths to candidate profiles
            position_title (str): Title of the position
        
        Returns:
            list: List of dictionaries containing candidate information and their emails
        """
        notifications = []
        
        for result in ranked_results:
            # Find the full profile path for this candidate
            profile_path = next(
                (path for path in profile_paths if os.path.basename(path) == result['filename']),
                None
            )
            
            if profile_path:
                candidate_info = {
                    'profile_path': profile_path,
                    'similarity_score': result['similarity_score']
                }
                
                # Generate email
                email_content = self.compose_shortlist_email(candidate_info, position_title)
                
                notifications.append({
                    'candidate_name': self.extract_candidate_name(profile_path),
                    'email_content': email_content,
                    'similarity_score': result['similarity_score']
                })
        
        return notifications
